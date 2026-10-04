import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette: Premium Ink-Blue / Instrument Theme
    BG_DARK = RGBColor(10, 16, 29)        # #0A101D
    CARD_BG = RGBColor(18, 28, 52)        # #121C34
    CARD_BORDER = RGBColor(30, 48, 86)    # #1E3056
    ACCENT_CYAN = RGBColor(56, 189, 248)  # #38BDF8
    ACCENT_AMBER = RGBColor(245, 158, 11) # #F59E0B
    ACCENT_EMERALD = RGBColor(16, 185, 129)# #10B981
    ACCENT_ROSE = RGBColor(244, 63, 94)   # #F43F5E
    TEXT_WHITE = RGBColor(248, 250, 252)  # #F8FAFC
    TEXT_MUTED = RGBColor(160, 175, 200)  # #A0AFC8
    TEXT_DIM = RGBColor(100, 116, 139)    # #64748B

    def set_slide_background(slide, color=BG_DARK):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, category, title, page_num):
        # Category / Eyebrow
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(10), Inches(0.35))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.font.name = "Arial"

        # Main Title
        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.6))
        tf2 = tx_box2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.font.name = "Arial"

        # Page Number / Footer
        tx_box3 = slide.shapes.add_textbox(Inches(10.5), Inches(7.0), Inches(2.0), Inches(0.3))
        tf3 = tx_box3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = f"DIGITAL TWIN — {page_num:02d}"
        p3.alignment = PP_ALIGN.RIGHT
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_DIM
        p3.font.name = "Arial"

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, RGBColor(7, 12, 24))

    # Left Hero Accent Bar
    hero_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.4), prs.slide_height)
    hero_bar.fill.solid()
    hero_bar.fill.fore_color.rgb = ACCENT_CYAN
    hero_bar.line.fill.background()

    # Title content
    tbox = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(4.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "STUDY AI AGENT // MINI PROJECT"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p2 = tf.add_paragraph()
    p2.text = "Digital Twin"
    p2.font.size = Pt(54)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    p3 = tf.add_paragraph()
    p3.text = "Autonomous AI Study Agent for Java DSA & Technical Interview Preparation"
    p3.font.size = Pt(22)
    p3.font.color.rgb = ACCENT_AMBER
    p3.space_before = Pt(8)

    p4 = tf.add_paragraph()
    p4.text = (
        "Ingests past quiz history to model learner mastery via mathematical recency-weighted decay, "
        "mines live LeetCode Java coding patterns, and synthesizes a 7-day adaptive revision runway "
        "with 0% duplicate practice challenges."
    )
    p4.font.size = Pt(15)
    p4.font.color.rgb = TEXT_MUTED
    p4.space_before = Pt(16)

    # Key Tag Chips
    tags = ["Recency-Decay Math (tau = 7d)", "LeetCode Java Scraping", "FastAPI Async", "Pydantic Guardrail", "Smoothed Terrain UI"]
    chip_left = Inches(1.2)
    for tag in tags:
        chip = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, chip_left, Inches(5.8), Inches(2.2), Inches(0.45))
        chip.fill.solid()
        chip.fill.fore_color.rgb = CARD_BG
        chip.line.color.rgb = CARD_BORDER
        tf_c = chip.text_frame
        p_c = tf_c.paragraphs[0]
        p_c.text = tag
        p_c.alignment = PP_ALIGN.CENTER
        p_c.font.size = Pt(10)
        p_c.font.color.rgb = ACCENT_CYAN
        chip_left += Inches(2.3)

    # =========================================================================
    # SLIDE 2: PROBLEM STATEMENT, OBJECTIVE & SCOPE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Project Foundation", "Problem Statement, Objective & Scope", 2)

    # Problem Card (Top Left)
    p_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(2.2))
    p_card.fill.solid()
    p_card.fill.fore_color.rgb = CARD_BG
    p_card.line.color.rgb = ACCENT_ROSE
    tf_p = p_card.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "01 | Problem Statement"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE
    p2 = tf_p.add_paragraph()
    p2.text = (
        "Students practice hundreds of coding questions but cannot convert past attempts into targeted action. "
        "Critical weak topics stay hidden behind static percentage scores, leading to generic, unoptimized revision "
        "instead of drilling actual gaps."
    )
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    # Objective Card (Top Right)
    o_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(2.2))
    o_card.fill.solid()
    o_card.fill.fore_color.rgb = CARD_BG
    o_card.line.color.rgb = ACCENT_EMERALD
    tf_o = o_card.text_frame
    tf_o.word_wrap = True
    p = tf_o.paragraphs[0]
    p.text = "02 | Core Objective"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p2 = tf_o.add_paragraph()
    p2.text = (
        "Design an autonomous Digital Twin study agent that ingests granular attempt history, detects weak topics via "
        "exponential recency decay, and auto-synthesizes a 7-day revision runway with authentic Java LeetCode practice questions."
    )
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_before = Pt(8)

    # Scope of Project (Bottom Wide Card)
    s_card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.9), Inches(11.7), Inches(2.8))
    s_card.fill.solid()
    s_card.fill.fore_color.rgb = CARD_BG
    s_card.line.color.rgb = CARD_BORDER
    tf_s = s_card.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "Scope of the Project (Technical Deliverables)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    scopes = [
        ("1", "Source-of-Truth Ingestion:", "Ingest student quiz history (topic, question, is_correct, time_spent, timestamp)."),
        ("2", "Recency-Weighted Decay Analytics:", "Derive mastery via exponential decay (tau = 7 days) and failure streak penalties."),
        ("3", "LeetCode Java Web Scraping:", "Mine authentic LeetCode problem patterns, Java starter code, constraints, and test cases."),
        ("4", "Deterministic AI Revision Runway:", "Auto-generate a 7-day study plan allocating ~70% study time to ranked weak topics."),
        ("5", "0% Duplicate Practice Guardrail:", "Cryptographic SHA-256 deduplication ensuring no duplicate problems per student per week.")
    ]

    for num, title_s, desc_s in scopes:
        p_item = tf_s.add_paragraph()
        p_item.text = f"[{num}] {title_s} {desc_s}"
        p_item.font.size = Pt(11.5)
        p_item.font.color.rgb = TEXT_WHITE
        p_item.space_before = Pt(5)

    # =========================================================================
    # SLIDE 3: HARDWARE & SOFTWARE REQUIREMENTS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "System Specifications", "Hardware & Software Requirements", 3)

    # Hardware Card (Left)
    hw_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    hw_card.fill.solid()
    hw_card.fill.fore_color.rgb = CARD_BG
    hw_card.line.color.rgb = CARD_BORDER
    tf_hw = hw_card.text_frame
    tf_hw.word_wrap = True
    p = tf_hw.paragraphs[0]
    p.text = "HARDWARE REQUIREMENTS"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    hw_items = [
        ("Processor", "Intel Core i5 / AMD Ryzen 5 or higher"),
        ("RAM", "8 GB minimum (16 GB recommended for local testing)"),
        ("Storage", "256 GB SSD (for DB storage, scraper cache & logs)"),
        ("GPU Support", "Optional — required only for local open-source LLM inference (Ollama/vLLM)"),
        ("Network", "Broadband internet connection for LeetCode GraphQL & LLM API calls"),
        ("Display", "1080p Full HD monitor for high-density instrument telemetry dashboard")
    ]
    for label, val in hw_items:
        p_i = tf_hw.add_paragraph()
        p_i.text = f"* {label}: {val}"
        p_i.font.size = Pt(12)
        p_i.font.color.rgb = TEXT_MUTED
        p_i.space_before = Pt(10)

    # Software Card (Right)
    sw_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    sw_card.fill.solid()
    sw_card.fill.fore_color.rgb = CARD_BG
    sw_card.line.color.rgb = CARD_BORDER
    tf_sw = sw_card.text_frame
    tf_sw.word_wrap = True
    p = tf_sw.paragraphs[0]
    p.text = "SOFTWARE & LIBRARIES STACK"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    sw_items = [
        ("Core Language", "Python 3.10+ (FastAPI async engine) & JavaScript (ES6+)"),
        ("Web Scraping & Mining", "BeautifulSoup4, httpx.AsyncClient, LeetCode Public GraphQL"),
        ("AI Agent & Validation", "Pydantic v2 (Strict Schema Contract), Gemini / OpenAI API"),
        ("Backend Framework", "FastAPI (async non-blocking routes, background job polling)"),
        ("Data Persistence", "Async SQLite (aiosqlite) with MySQL & MongoDB schemas"),
        ("Frontend UI / UX", "React + Vite, Bespoke Instrument CSS (No generic Tailwind)"),
        ("Testing & Tooling", "PyTest, PyTest-Asyncio, Git/GitHub, Postman, VS Code")
    ]
    for label, val in sw_items:
        p_i = tf_sw.add_paragraph()
        p_i.text = f"* {label}: {val}"
        p_i.font.size = Pt(12)
        p_i.font.color.rgb = TEXT_MUTED
        p_i.space_before = Pt(10)

    # =========================================================================
    # SLIDE 4: ARCHITECTURE & 8-WEEK TIMELINE (TWO HALVES)
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Engineering Roadmap", "Architecture, Process Flow & Timeline", 4)

    # Process Flow (6 Steps)
    steps = [
        ("1. Ingestion", "Quiz attempt history logged to DB"),
        ("2. Recency Decay", "tau=7d decay & weakness ranking"),
        ("3. LeetCode Scrape", "Mine Java question patterns"),
        ("4. Deterministic AI", "Pydantic-validated 7-day plan"),
        ("5. Deduplication", "0% duplicate questions check"),
        ("6. Instrument UI", "Readiness dial & terrain curve")
    ]

    step_w = Inches(1.8)
    step_l = Inches(0.8)
    for title_st, desc_st in steps:
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, step_l, Inches(1.5), step_w, Inches(1.2))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_CYAN
        tf_b = box.text_frame
        tf_b.word_wrap = True
        p = tf_b.paragraphs[0]
        p.text = title_st
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p2 = tf_b.add_paragraph()
        p2.text = desc_st
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED
        step_l += Inches(1.95)

    # Timeline Section Header
    tbox = s4.shapes.add_textbox(Inches(0.8), Inches(2.9), Inches(11.7), Inches(0.4))
    tf_t = tbox.text_frame
    p = tf_t.paragraphs[0]
    p.text = "PROJECT TIMELINE (8 WEEKS IN TWO INDEPENDENTLY-DELIVERABLE HALVES)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER

    # Half 1 Box
    h1_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.4), Inches(5.6), Inches(3.4))
    h1_box.fill.solid()
    h1_box.fill.fore_color.rgb = CARD_BG
    h1_box.line.color.rgb = ACCENT_CYAN
    tf_h1 = h1_box.text_frame
    tf_h1.word_wrap = True
    p = tf_h1.paragraphs[0]
    p.text = "DELIVERABLE HALF 1 (WEEKS 1 - 4)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p_sub = tf_h1.add_paragraph()
    p_sub.text = "Core Analytics, LeetCode Scraper & Instrument Panel"
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = TEXT_DIM

    h1_pts = [
        ("Week 1", "Data schema & quiz attempts source-of-truth ingestion engine."),
        ("Week 2", "LeetCode Java question scraper & offline resilient pattern catalog."),
        ("Week 3", "FastAPI /topics/analyze with recency-weighted decay mathematics."),
        ("Week 4", "Dark instrument dashboard with signature smoothed terrain curve.")
    ]
    for wk, desc in h1_pts:
        p_i = tf_h1.add_paragraph()
        p_i.text = f"* {wk}: {desc}"
        p_i.font.size = Pt(11)
        p_i.font.color.rgb = TEXT_WHITE
        p_i.space_before = Pt(6)

    # Half 2 Box
    h2_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(3.4), Inches(5.7), Inches(3.4))
    h2_box.fill.solid()
    h2_box.fill.fore_color.rgb = CARD_BG
    h2_box.line.color.rgb = ACCENT_EMERALD
    tf_h2 = h2_box.text_frame
    tf_h2.word_wrap = True
    p = tf_h2.paragraphs[0]
    p.text = "DELIVERABLE HALF 2 (WEEKS 5 - 8)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    p_sub = tf_h2.add_paragraph()
    p_sub.text = "Deterministic AI Agent, Runway & Verification Suite"
    p_sub.font.size = Pt(10.5)
    p_sub.font.color.rgb = TEXT_DIM

    h2_pts = [
        ("Week 5", "Deterministic LLM planner with Pydantic validation & retry loops."),
        ("Week 6", "7-Day revision runway with 70% dynamic weak-topic allocation."),
        ("Week 7", "Non-duplicate practice question drawer with Java code runner."),
        ("Week 8", "Async job polling (>2s), latency SLA (<300ms) & automated tests.")
    ]
    for wk, desc in h2_pts:
        p_i = tf_h2.add_paragraph()
        p_i.text = f"* {wk}: {desc}"
        p_i.font.size = Pt(11)
        p_i.font.color.rgb = TEXT_WHITE
        p_i.space_before = Pt(6)

    # =========================================================================
    # SLIDE 5: LEETCODE JAVA SCRAPING & QUESTION PIPELINE (NEW SLIDE)
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Domain-Specific Question Mining", "LeetCode Java Web Scraping & Practice Pipeline", 5)

    col1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(3.6), Inches(5.2))
    col1.fill.solid()
    col1.fill.fore_color.rgb = CARD_BG
    col1.line.color.rgb = ACCENT_CYAN
    tf_c1 = col1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "1. Live GraphQL Scraper"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    c1_items = [
        "Connects to LeetCode public GraphQL endpoints using httpx.AsyncClient.",
        "Scrapes authentic question metadata: title, slug, difficulty (Easy/Med/Hard), and category tags.",
        "Covers 8 core DSA categories: Arrays, Two Pointers, Sliding Window, Binary Search, Trees, Graphs, DP, Heaps.",
        "Automatic fallback to curated offline catalog during rate limits or Cloudflare blocks."
    ]
    for it in c1_items:
        p_it = tf_c1.add_paragraph()
        p_it.text = f"* {it}"
        p_it.font.size = Pt(11)
        p_it.font.color.rgb = TEXT_MUTED
        p_it.space_before = Pt(8)

    col2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.8), Inches(1.5), Inches(3.7), Inches(5.2))
    col2.fill.solid()
    col2.fill.fore_color.rgb = CARD_BG
    col2.line.color.rgb = ACCENT_AMBER
    tf_c2 = col2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "2. Java Template Extraction"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_AMBER
    c2_items = [
        "Extracts standard Java method signatures: class Solution { public int[] ... }.",
        "Includes sample test cases with input/output assertions for automated verification.",
        "Includes JVM performance tips (e.g. Arrays.sort() dual-pivot quicksort, HashMap capacity sizing).",
        "Flags Java edge cases (Integer object cache -128..127, 32-bit signed overflow prevention)."
    ]
    for it in c2_items:
        p_it = tf_c2.add_paragraph()
        p_it.text = f"* {it}"
        p_it.font.size = Pt(11)
        p_it.font.color.rgb = TEXT_MUTED
        p_it.space_before = Pt(8)

    col3 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.5), Inches(3.6), Inches(5.2))
    col3.fill.solid()
    col3.fill.fore_color.rgb = CARD_BG
    col3.line.color.rgb = ACCENT_EMERALD
    tf_c3 = col3.text_frame
    tf_c3.word_wrap = True
    p = tf_c3.paragraphs[0]
    p.text = "3. 0% Duplicate Guardrail"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_EMERALD
    c3_items = [
        "Cryptographic SHA-256 fingerprinting on topic and normalized question signatures.",
        "Maintains persistent question index in SQLite/MySQL per student.",
        "Strictly guarantees 0% duplicate questions per topic across weekly revision plans.",
        "If all catalog items were recently practiced, generates unique target drill variants automatically."
    ]
    for it in c3_items:
        p_it = tf_c3.add_paragraph()
        p_it.text = f"* {it}"
        p_it.font.size = Pt(11)
        p_it.font.color.rgb = TEXT_MUTED
        p_it.space_before = Pt(8)

    # =========================================================================
    # SLIDE 6: MATHEMATICAL MODEL & MASTERY TERRAIN (NEW SLIDE)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Core Mathematics & Signature UI", "Recency-Weighted Accuracy & Mastery Terrain", 6)

    # Left: Mathematical Formulas
    m_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = CARD_BG
    m_box.line.color.rgb = ACCENT_CYAN
    tf_m = m_box.text_frame
    tf_m.word_wrap = True
    p = tf_m.paragraphs[0]
    p.text = "MATHEMATICAL RECENCY DECAY MODEL"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    math_desc = [
        ("1. Exponential Decay Weighting", "w_i = exp(-lambda * (t_now - t_i)), where lambda = ln(2) / tau with half-life tau = 7 days (604,800s). Recent attempts carry dramatically higher influence than old successes."),
        ("2. Recency-Weighted Accuracy", "Acc_w(T) = sum(w_i * is_correct_i) / sum(w_i). Unifies historical mastery with immediate recall state."),
        ("3. Streak Penalty & Urgency", "Weakness = (1.0 - Acc_w) * (1.0 + 0.20 * min(streak, 4)). Penalizes consecutive recent failures immediately into CRITICAL status regardless of ancient quiz scores."),
        ("4. Readiness Telemetry (0 - 100%)", "Curriculum-wide readiness score computed across topics with historical trend delta (+/- vs previous checkpoint).")
    ]
    for heading, txt in math_desc:
        p_h = tf_m.add_paragraph()
        p_h.text = heading
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = ACCENT_AMBER
        p_h.space_before = Pt(8)
        p_t = tf_m.add_paragraph()
        p_t.text = txt
        p_t.font.size = Pt(10.5)
        p_t.font.color.rgb = TEXT_MUTED

    # Right: Signature Visual Moment
    v_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    v_box.fill.solid()
    v_box.fill.fore_color.rgb = CARD_BG
    v_box.line.color.rgb = ACCENT_ROSE
    tf_v = v_box.text_frame
    tf_v.word_wrap = True
    p = tf_v.paragraphs[0]
    p.text = "SIGNATURE VISUAL MOMENT: WEAKNESS TERRAIN"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ROSE

    vis_desc = [
        ("Topological Elevation vs Generic Charts", "Replaces standard SaaS bar/pie charts with an organic, smoothed terrain curve plotted using SVG cubic Bezier spline interpolation."),
        ("Summit Elevation = Weakness Intensity", "High mountain peaks (>75m elevation) represent urgent challenge zones (e.g. Dynamic Programming); mastered domains form conquered valleys and plains (<25m)."),
        ("Interactive Telemetry Beacons", "Glowing anchor nodes with radar ripple effects on hover/focus, revealing real-time accuracy %, failure streaks, and attempt counts."),
        ("Zero Cumulative Layout Shift (CLS)", "Fixed aspect-ratio containers with shimmering precision skeletons ensure zero layout shift during async backend loading.")
    ]
    for heading, txt in vis_desc:
        p_h = tf_v.add_paragraph()
        p_h.text = heading
        p_h.font.size = Pt(12)
        p_h.font.bold = True
        p_h.font.color.rgb = TEXT_WHITE
        p_h.space_before = Pt(8)
        p_t = tf_v.add_paragraph()
        p_t.text = txt
        p_t.font.size = Pt(10.5)
        p_t.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 7: USABILITY & APPLICATIONS
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Real-World Impact", "Usability & Target Applications", 7)

    use_cases = [
        ("S", "Software Engineers & Tech Aspirants", "Targeted DSA preparation for students and professionals grinding LeetCode in Java for top-tier software engineering interviews.", ACCENT_CYAN),
        ("C", "Competitive Programming Aspirants", "Eliminates blind spots across complex graph algorithms, dynamic programming transitions, and monotonic stacks.", ACCENT_AMBER),
        ("O", "Online EdTech & Bootcamp Platforms", "Drop-in adaptive intelligence module that turns static quiz platforms into personalized learning twins.", ACCENT_EMERALD),
        ("T", "Universities & CS Instructors", "Class-level telemetry surfaces syllabus topics where students struggle most, enabling data-driven remedial lectures.", ACCENT_ROSE),
        ("L", "Autonomous Self-Learners", "Provides a structured, dynamically evolving 7-day revision trajectory without requiring an expensive private tutor.", ACCENT_CYAN),
        ("R", "Enterprise & Certification Training", "Enables technical companies to train enterprise software teams on modern Java backend patterns with targeted drills.", ACCENT_AMBER),
    ]

    card_coords = [
        (Inches(0.8), Inches(1.5)), (Inches(4.8), Inches(1.5)), (Inches(8.8), Inches(1.5)),
        (Inches(0.8), Inches(4.1)), (Inches(4.8), Inches(4.1)), (Inches(8.8), Inches(4.1))
    ]

    for idx, (badge, title_u, desc_u, color_u) in enumerate(use_cases):
        x, y = card_coords[idx]
        card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.4))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        tf_c = card.text_frame
        tf_c.word_wrap = True
        p = tf_c.paragraphs[0]
        p.text = f"[{badge}]  {title_u}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = color_u
        p2 = tf_c.add_paragraph()
        p2.text = desc_u
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(8)

    # =========================================================================
    # SLIDE 8: CONTRIBUTION OF EACH MEMBER
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Team Execution", "Contribution of Each Member", 8)

    # Table Layout
    table_shape = s8.shapes.add_table(7, 3, Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(3.2)
    table.columns[2].width = Inches(6.5)

    headers_t = ["Team Member", "Role", "Key Engineering Contribution"]
    for i, h in enumerate(headers_t):
        cell = table.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(22, 36, 66)
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN

    members_data = [
        ("Member 1", "Team Lead / Backend Developer", "Agent architecture, FastAPI async routes, non-blocking job manager & recency decay analytics."),
        ("Member 2", "Data & Scraper Engineer", "LeetCode Java web scraper & GraphQL pipeline, problem catalog curation & database source of truth."),
        ("Member 3", "AI & Planning Engineer", "Deterministic 7-day study planner, Pydantic strict schema contract & 0% duplicate guardrail."),
        ("Member 4", "Frontend UI / UX Developer", "Instrument-panel design system, smoothed terrain curve (SVG Bézier) & Java practice question drawer."),
        ("Member 5", "Testing & Verification Lead", "Automated PyTest suite (recency decay math, p95 <300ms SLA, Pydantic validation & zero duplicate check)."),
        ("Member 6", "Research & DevOps Engineer", "Literature grounding (Ebbinghaus decay), cloud deployment readiness & demonstration benchmarking.")
    ]

    for row_idx, (m, r, c) in enumerate(members_data, start=1):
        for col_idx, text_val in enumerate([m, r, c]):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            p = cell.text_frame.paragraphs[0]
            p.text = text_val
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_WHITE if col_idx < 2 else TEXT_MUTED
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = ACCENT_AMBER

    # =========================================================================
    # SLIDE 9: REFERENCES
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Academic & Technical Foundations", "References & Documentation", 9)

    refs = [
        ("1", "LeetCode API & Problemset Taxonomy", "Public GraphQL schema and algorithmic coding taxonomy for Java DSA problems (leetcode.com)."),
        ("2", "Pydantic v2 Documentation", "Server-side data validation, deterministic schema enforcement, and JSON contract validation (docs.pydantic.dev)."),
        ("3", "FastAPI Async Web Framework", "High-performance async Python backend architecture and Starlette event loop specifications (fastapi.tiangolo.com)."),
        ("4", "Ebbinghaus, H. — Memory & The Forgetting Curve", "Memory: A Contribution to Experimental Psychology. Theoretical basis for the 7-day exponential recency decay model."),
        ("5", "BeautifulSoup4 & HTTPX", "Modern async web scraping and HTTP client architectures for Python (curl / GraphQL scraping)."),
        ("6", "IEEE Xplore Research Papers", "Adaptive learning systems, knowledge tracing, and personalised e-learning algorithms for computer science students.")
    ]

    ref_y = Inches(1.5)
    for num, title_r, desc_r in refs:
        r_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), ref_y, Inches(11.7), Inches(0.75))
        r_card.fill.solid()
        r_card.fill.fore_color.rgb = CARD_BG
        r_card.line.color.rgb = CARD_BORDER
        tf_r = r_card.text_frame
        tf_r.word_wrap = True
        p = tf_r.paragraphs[0]
        p.text = f"[{num}]  {title_r} — {desc_r}"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_WHITE
        ref_y += Inches(0.85)

    output_path = "Digital_Twin_AI_Study_Agent_Presentation.pptx"
    prs.save(output_path)
    print(f"[Success] Generated PowerPoint presentation at: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_presentation()
